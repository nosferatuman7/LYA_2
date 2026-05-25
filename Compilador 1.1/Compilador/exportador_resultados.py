# -*- coding: utf-8 -*-
"""
exportador_resultados.py
========================
EXPORTACION DE RESULTADOS.

Exporta los resultados del compilador a un archivo para revisarlos fuera de la
aplicacion. Si la libreria 'openpyxl' esta disponible, genera un libro de Excel
(.xlsx) con una hoja por cada resultado; si no, genera un reporte de texto (.txt)
equivalente, de modo que la exportacion siempre funcione.

Hojas / secciones generadas:
  - Tokens                 (tipo, valor, linea, columna)
  - Palabras reservadas    (palabra, cantidad de veces usada)
  - Tabla de simbolos      (nombre, tipo, clase, valor, linea, columna)
  - Codigo intermedio      (TAC)
  - Codigo intermedio optimizado
  - Codigo objeto          (Python / Raspberry Pi)
"""


class ExportadorResultados:
    """
    Uso:
        ruta = ExportadorResultados().exportar(ruta_base, datos)

    'ruta_base' es la ruta SIN extension. 'datos' es un diccionario con:
        tokens          -> lista de tuplas (tipo, valor, linea, columna)
        palabras        -> lista de tuplas (palabra, cantidad)
        simbolos        -> lista de dicts con datos del simbolo
        intermedio      -> str (codigo intermedio)
        intermedio_opt  -> str (codigo intermedio optimizado)
        objeto          -> str (codigo objeto)
    """

    def exportar(self, ruta_base, datos):
        try:
            import openpyxl  # noqa: F401
            ruta = ruta_base + ".xlsx"
            self._exportar_xlsx(ruta, datos)
            return ruta
        except ImportError:
            ruta = ruta_base + ".txt"
            self._exportar_txt(ruta, datos)
            return ruta

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------

    def _exportar_xlsx(self, ruta, datos):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()

        # --- Hoja: Tokens ---
        hoja = wb.active
        hoja.title = "Tokens"
        self._encabezado(hoja, ["Tipo", "Valor", "Linea", "Columna"], Font)
        for (tipo, valor, linea, col) in datos.get("tokens", []):
            hoja.append([tipo, valor, linea, col])
        self._ajustar(hoja)

        # --- Hoja: Palabras reservadas ---
        hoja = wb.create_sheet("Palabras reservadas")
        self._encabezado(hoja, ["Palabra", "Cantidad"], Font)
        for (palabra, cantidad) in datos.get("palabras", []):
            hoja.append([palabra, cantidad])
        self._ajustar(hoja)

        # --- Hoja: Tabla de simbolos ---
        hoja = wb.create_sheet("Tabla de simbolos")
        self._encabezado(hoja, ["Nombre", "Tipo", "Clase", "Valor", "Linea", "Columna"], Font)
        for s in datos.get("simbolos", []):
            hoja.append([
                s.get("nombre"), s.get("tipo"), s.get("clase"),
                self._texto_valor(s.get("valor")), s.get("linea"), s.get("columna"),
            ])
        self._ajustar(hoja)

        # --- Hoja: Codigo intermedio ---
        self._hoja_texto(wb, "Codigo intermedio", datos.get("intermedio", ""), Font)

        # --- Hoja: Codigo intermedio optimizado ---
        self._hoja_texto(wb, "Codigo intermedio optimizado", datos.get("intermedio_opt", ""), Font)

        # --- Hoja: Codigo objeto ---
        self._hoja_texto(wb, "Codigo objeto", datos.get("objeto", ""), Font)

        wb.save(ruta)

    def _encabezado(self, hoja, columnas, Font):
        hoja.append(columnas)
        for celda in hoja[1]:
            celda.font = Font(bold=True)

    def _hoja_texto(self, wb, titulo, texto, Font):
        hoja = wb.create_sheet(titulo[:31])  # Excel limita el nombre a 31 chars
        for linea in (texto or "").split("\n"):
            hoja.append([linea])
        try:
            hoja.column_dimensions["A"].width = 70
        except Exception:
            pass

    def _ajustar(self, hoja):
        for columna in hoja.columns:
            ancho = 10
            letra = columna[0].column_letter
            for celda in columna:
                if celda.value is not None:
                    ancho = max(ancho, len(str(celda.value)) + 2)
            hoja.column_dimensions[letra].width = min(ancho, 60)

    # ------------------------------------------------------------------
    # Texto (respaldo si no hay openpyxl)
    # ------------------------------------------------------------------

    def _exportar_txt(self, ruta, datos):
        lineas = []
        lineas.append("=" * 60)
        lineas.append("RESULTADOS DEL COMPILADOR")
        lineas.append("=" * 60)

        lineas.append("\n=== TOKENS ===")
        lineas.append(f"{'TIPO':<22}{'VALOR':<22}{'LINEA':<8}{'COLUMNA'}")
        for (tipo, valor, linea, col) in datos.get("tokens", []):
            lineas.append(f"{str(tipo):<22}{str(valor):<22}{str(linea):<8}{str(col)}")

        lineas.append("\n=== PALABRAS RESERVADAS (conteo) ===")
        lineas.append(f"{'PALABRA':<22}{'CANTIDAD'}")
        for (palabra, cantidad) in datos.get("palabras", []):
            lineas.append(f"{str(palabra):<22}{str(cantidad)}")

        lineas.append("\n=== TABLA DE SIMBOLOS ===")
        for s in datos.get("simbolos", []):
            lineas.append(
                f"nombre: {s.get('nombre')}, tipo: {s.get('tipo')}, "
                f"clase: {s.get('clase')}, valor: {self._texto_valor(s.get('valor'))}, "
                f"linea: {s.get('linea')}, columna: {s.get('columna')}"
            )

        lineas.append("\n=== CODIGO INTERMEDIO ===")
        lineas.append(datos.get("intermedio", ""))

        lineas.append("\n=== CODIGO INTERMEDIO OPTIMIZADO ===")
        lineas.append(datos.get("intermedio_opt", ""))

        lineas.append("\n=== CODIGO OBJETO (Python / Raspberry Pi) ===")
        lineas.append(datos.get("objeto", ""))

        with open(ruta, "w", encoding="utf-8") as f:
            f.write("\n".join(lineas))

    # ------------------------------------------------------------------
    # Utilidades
    # ------------------------------------------------------------------

    @staticmethod
    def _texto_valor(valor):
        if valor is None:
            return ""
        if isinstance(valor, list):
            return ", ".join(str(v) for v in valor)
        return str(valor)


# Compatibilidad con el estilo del proyecto
exportador_resultados = ExportadorResultados
